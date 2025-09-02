import os
import json
import uuid
import shutil
import logging
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import JSONResponse, FileResponse, HTMLResponse
from dotenv import load_dotenv
import PyPDF2
import docx

from langchain.chat_models import ChatOpenAI
from langchain.prompts import PromptTemplate
from langchain import LLMChain

# Logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Load environment variables
load_dotenv()
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")

app = FastAPI(title="LangChain Compliance Agent with Modify Feature")
UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

ALLOWED_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/msword": ".doc"
}

def extract_text_from_pdf(file_path: Path) -> str:
    text = ""
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
        return text.strip()
    except Exception as e:
        logger.exception("PDF extraction error")
        raise

def extract_text_from_docx(file_path: Path) -> str:
    try:
        document = docx.Document(file_path)
        return "\n".join([p.text for p in document.paragraphs if p.text]).strip()
    except Exception as e:
        logger.exception("DOCX extraction error")
        raise

def save_to_docx(text: str) -> Path:
    filename = f"modified_{uuid.uuid4()}.docx"
    output_path = UPLOAD_DIR / filename
    try:
        doc = docx.Document()
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(output_path)
        return output_path
    except Exception as e:
        logger.exception("DOCX save error")
        raise

def analyze_and_modify(text: str, guidelines: str) -> dict:
    if not OPENAI_API_KEY:
        raise RuntimeError("Missing OPENAI_API_KEY in .env")
    llm = ChatOpenAI(model_name="gpt-4o-mini", temperature=0, openai_api_key=OPENAI_API_KEY)

    prompt = PromptTemplate(
        input_variables=["guidelines", "text"],
        template="""
        You are an expert editor. Analyze the following DOCUMENT against the GUIDELINES
        and rewrite it to be compliant. Return JSON with:
        - report: Compliance summary and issues
        - modified_text: Fully rewritten document text

        GUIDELINES:
        {guidelines}

        DOCUMENT:
        {text}
        """
    )
    chain = LLMChain(llm=llm, prompt=prompt)

    try:
        raw_response = chain.predict(guidelines=guidelines, text=text[:5000])
        return json.loads(raw_response)
    except json.JSONDecodeError:
        logger.warning("Failed to parse JSON from LLM. Returning raw response.")
        return {
            "report": {"summary": "Could not parse JSON", "raw_output": raw_response},
            "modified_text": raw_response
        }
    except Exception:
        logger.exception("LLM call failed")
        raise


@app.get("/", response_class=HTMLResponse)
async def home():
    """Simple home page with upload form."""
    return """
    <html>
      <head>
        <title>Compliance Checker</title>
      </head>
      <body style="font-family: Arial; margin: 40px;">
        <h1>Upload Document for Analysis & Modification</h1>
        <form action="/process/" enctype="multipart/form-data" method="post">
          <label for="file">Choose file:</label><br>
          <input type="file" name="file" accept=".pdf,.doc,.docx" required><br><br>

          <label for="guidelines">Guidelines:</label><br>
          <textarea name="guidelines" rows="5" cols="60"
            placeholder="Enter writing guidelines here..." required></textarea><br><br>

          <button type="submit">Upload & Process</button>
        </form>
      </body>
    </html>
    """

@app.post("/process/")
async def process_file(file: UploadFile = File(...), guidelines: str = Form(...)):
    # Validate
    if file.content_type not in ALLOWED_TYPES:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {file.content_type}")

    ext = ALLOWED_TYPES[file.content_type]
    filename = f"{uuid.uuid4()}{ext}"
    file_path = UPLOAD_DIR / filename

    # Save upload
    try:
        with file_path.open("wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception:
        logger.exception("File save error")
        raise HTTPException(status_code=500, detail="File save error")

    # Extract
    try:
        text = extract_text_from_pdf(file_path) if ext == ".pdf" else extract_text_from_docx(file_path)
        if not text.strip():
            raise ValueError("No extractable text")
    except Exception:
        logger.exception("Text extraction failed")
        raise HTTPException(status_code=500, detail="Text extraction failed")

    # Analyze + rewrite
    try:
        result = analyze_and_modify(text, guidelines)
    except Exception:
        logger.exception("LLM analysis failed")
        raise HTTPException(status_code=500, detail="LLM analysis failed")

    # Save rewritten
    rewritten_text = result.get("modified_text", "")
    try:
        rewritten_file = save_to_docx(rewritten_text)
    except Exception:
        logger.exception("File save failed")
        raise HTTPException(status_code=500, detail="Could not save rewritten file")

    return {
        "report": result.get("report", {}),
        "download_link": f"/download/{rewritten_file.name}"
    }

@app.get("/download/{filename}")
async def download_file(filename: str):
    file_path = UPLOAD_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(file_path, filename=filename)
